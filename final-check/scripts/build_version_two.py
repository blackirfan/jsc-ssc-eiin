import os, re, glob, time, importlib.util
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # final-check/
PROJECT_ROOT = os.path.normpath(os.path.join(ROOT, ".."))           # jsc-ssc-eiin/
PDF_DIR = os.path.join(PROJECT_ROOT, "pdfs")                        # authoritative, most up to date

spec = importlib.util.spec_from_file_location("audit", os.path.join(ROOT, "scripts", "audit.py"))
audit = importlib.util.module_from_spec(spec)
spec.loader.exec_module(audit)

ABS_RE = re.compile(r'^ABS\.?$', re.IGNORECASE)

SSC_YEARS = list(range(2011, 2027))
JSC_YEARS = list(range(2011, 2020))

NAME_RE = re.compile(r'Institution:\s+(.+?)\s+\(EIIN:')


def summarize_v2(entries):
    total = len(entries)
    passed = 0
    absent = 0
    gpa5 = 0
    known_sum = 0.0
    unknown_pass_count = 0
    for roll, content in entries:
        if content == '':
            passed += 1
            unknown_pass_count += 1
        elif audit.NUM_RE.match(content):
            val = float(content)
            passed += 1
            known_sum += val
            if abs(val - 5.0) < 1e-9:
                gpa5 += 1
        elif ABS_RE.match(content):
            absent += 1
        else:
            pass  # F1/F2/F3/... failed but appeared
    appeared = total - absent
    return {
        'total': total, 'appeared': appeared, 'passed': passed,
        'gpa5': gpa5, 'known_sum': known_sum,
        'unknown_pass_count': unknown_pass_count,
    }


def compute_row(kind, eiin, year):
    path = os.path.join(PDF_DIR, f"{kind}_{eiin}_{year}.pdf")
    if not os.path.exists(path):
        return None
    parsed, err = audit.parse_pdf(path)
    if err:
        return None
    header = parsed['header']
    groups = parsed['groups']

    # institution name straight from the PDF text (more reliable per-year than
    # copying a single cached name across all years)
    name = None
    m = NAME_RE.search(parsed['raw_text'])
    if m:
        name = m.group(1).strip()

    science_entries = groups.get('SCIENCE', [])
    nonscience_entries = []
    for gname, entries in groups.items():
        if gname != 'SCIENCE':
            nonscience_entries.extend(entries)

    science_sum = summarize_v2(science_entries)
    nonscience_sum = summarize_v2(nonscience_entries)

    def sum_gpa_of(s):
        if s['unknown_pass_count'] > 0:
            return None
        return round(s['known_sum'], 2)

    result = {
        'school_name_pulled': name,
        'examinee_total': header['examinee'],
        'examinee_appeared': header['appeared'],
        'examinee_passed': header['passed'],
        'examinee_gpa5': header['gpa5'],
    }

    # overall sum_gpa = sum across all groups (None if any group has unknown GPA data)
    overall_summary = summarize_v2([e for entries in groups.values() for e in entries])
    result['sum_gpa'] = sum_gpa_of(overall_summary)

    if kind == 'ssc':
        result['science_examinee_total'] = science_sum['total'] if science_entries else None
        result['science_examinee_appeared'] = science_sum['appeared'] if science_entries else None
        result['science_examinee_passed'] = science_sum['passed'] if science_entries else None
        result['science_examinee_gpa5'] = science_sum['gpa5'] if science_entries else None
        result['science_sum_gpa'] = sum_gpa_of(science_sum) if science_entries else None

        result['nonscience_examinee_total'] = nonscience_sum['total'] if nonscience_entries else None
        result['nonscience_examinee_appeared'] = nonscience_sum['appeared'] if nonscience_entries else None
        result['nonscience_examinee_passed'] = nonscience_sum['passed'] if nonscience_entries else None
        result['nonscience_examinee_gpa5'] = nonscience_sum['gpa5'] if nonscience_entries else None
        result['nonscience_sum_gpa'] = sum_gpa_of(nonscience_sum) if nonscience_entries else None

    return result


def build(kind, template_ws, years):
    header = [c.value for c in template_ws[1]]
    eiin_col = header.index('eiin')
    exam_col = header.index('exam')

    schools = []
    for row in template_ws.iter_rows(min_row=2, values_only=True):
        if row[eiin_col] is None:
            continue
        schools.append((row[eiin_col], row[exam_col]))

    out_rows = []
    n_found, n_missing = 0, 0
    for eiin, exam in schools:
        for year in years:
            row = {h: None for h in header}
            row['eiin'] = eiin
            row['exam'] = exam
            row['exam_year'] = year
            data = compute_row(kind, eiin, year)
            if data is None:
                n_missing += 1
            else:
                n_found += 1
                for k, v in data.items():
                    if k in row:
                        row[k] = v
            out_rows.append(row)
    return header, out_rows, n_found, n_missing


def write_sheet(ws, header, rows):
    # clear existing data rows (keep header row 1)
    ws.delete_rows(2, ws.max_row)
    for r_i, row in enumerate(rows, start=2):
        for c_i, h in enumerate(header, start=1):
            ws.cell(row=r_i, column=c_i, value=row[h])


def main():
    target = os.path.join(ROOT, "excel", "ssc_jsc_results_scraped_version_two.xlsx")
    print("Loading template:", target)
    wb = openpyxl.load_workbook(target, data_only=True)

    t0 = time.time()
    header_ssc, rows_ssc, found_ssc, missing_ssc = build('ssc', wb['ssc'], SSC_YEARS)
    print(f"ssc: {len(rows_ssc)} rows built | pdf found={found_ssc} missing={missing_ssc} | {time.time()-t0:.1f}s")

    t0 = time.time()
    header_jsc, rows_jsc, found_jsc, missing_jsc = build('jsc', wb['jsc'], JSC_YEARS)
    print(f"jsc: {len(rows_jsc)} rows built | pdf found={found_jsc} missing={missing_jsc} | {time.time()-t0:.1f}s")

    write_sheet(wb['ssc'], header_ssc, rows_ssc)
    write_sheet(wb['jsc'], header_jsc, rows_jsc)

    wb.save(target)
    print("Saved:", target)


if __name__ == '__main__':
    main()
