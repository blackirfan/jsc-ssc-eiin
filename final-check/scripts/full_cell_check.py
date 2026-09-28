import os, glob, re, importlib.util, time
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

spec = importlib.util.spec_from_file_location("audit", os.path.join(ROOT, "scripts", "audit.py"))
audit = importlib.util.module_from_spec(spec)
spec.loader.exec_module(audit)

ALL_FIELDS = [
    'examinee_total', 'examinee_passed', 'examinee_gpa5', 'avg_gpa',
    'science_examinee_total', 'science_examinee_passed', 'science_examinee_gpa5', 'science_avg_gpa',
    'nonscience_examinee_total', 'nonscience_examinee_passed', 'nonscience_examinee_gpa5', 'nonscience_avg_gpa',
]

wb = openpyxl.load_workbook("excel/ssc_jsc_results_corrected.xlsx", data_only=True)

mismatches = []
unrecoverable_blanks = []
missing_pdf_rows = []

for kind in ('ssc', 'jsc'):
    ws = wb[kind]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    t0 = time.time()
    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx['scrape_status']] != 'done':
            continue
        n += 1
        eiin, year = row[idx['eiin']], row[idx['exam_year']]
        path = f"pdfs/{kind}_{eiin}_{year}.pdf"
        if not os.path.exists(path):
            missing_pdf_rows.append((kind, eiin, year))
            continue
        parsed, err = audit.parse_pdf(path)
        if err:
            missing_pdf_rows.append((kind, eiin, year, err))
            continue
        header_h = parsed['header']
        groups = parsed['groups']
        science_entries = groups.get('SCIENCE', [])
        nonscience_entries = []
        for gname, entries in groups.items():
            if gname != 'SCIENCE':
                nonscience_entries.extend(entries)
        all_entries = []
        for entries in groups.values():
            all_entries.extend(entries)
        overall_sum = audit.summarize(all_entries)
        science_sum = audit.summarize(science_entries)
        nonscience_sum = audit.summarize(nonscience_entries)
        overall_avg, _ = audit.avg_of(overall_sum)
        science_avg, _ = audit.avg_of(science_sum) if science_entries else (None, 'no_group')
        nonscience_avg, _ = audit.avg_of(nonscience_sum) if nonscience_entries else (None, 'no_group')
        science_gpa5 = None if science_sum['unknown_pass_count'] > 0 else (science_sum['gpa5'] if science_entries else None)
        nonscience_gpa5 = None if nonscience_sum['unknown_pass_count'] > 0 else (nonscience_sum['gpa5'] if nonscience_entries else None)

        computed = {
            'examinee_total': header_h['examinee'],
            'examinee_passed': header_h['passed'],
            'examinee_gpa5': header_h['gpa5'],
            'avg_gpa': overall_avg,
            'science_examinee_total': science_sum['total'] if science_entries else None,
            'science_examinee_passed': science_sum['passed'] if science_entries else None,
            'science_examinee_gpa5': science_gpa5,
            'science_avg_gpa': science_avg,
            'nonscience_examinee_total': nonscience_sum['total'] if nonscience_entries else None,
            'nonscience_examinee_passed': nonscience_sum['passed'] if nonscience_entries else None,
            'nonscience_examinee_gpa5': nonscience_gpa5,
            'nonscience_avg_gpa': nonscience_avg,
        }

        for field in ALL_FIELDS:
            excel_val = row[idx[field]]
            pdf_val = computed[field]
            if field.endswith('avg_gpa') and pdf_val is None:
                if excel_val is not None:
                    pass  # excel has a value pdf can't confirm/deny (shouldn't happen but note)
                continue
            cmp_excel, cmp_pdf = excel_val, pdf_val
            if not field.endswith('avg_gpa'):
                if cmp_excel is None:
                    cmp_excel = 0
                if cmp_pdf is None:
                    cmp_pdf = 0
            if not audit.close(cmp_excel, cmp_pdf):
                mismatches.append((kind, eiin, year, field, excel_val, pdf_val))

    print(f"{kind}: checked {n} done rows in {time.time()-t0:.1f}s")

print()
print("TOTAL mismatches across ALL fields:", len(mismatches))
for m in mismatches[:30]:
    print("  ", m)
print()
print("Rows with missing/unparseable PDF:", len(missing_pdf_rows))
for m in missing_pdf_rows:
    print("  ", m)
