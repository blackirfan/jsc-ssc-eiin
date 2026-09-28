import openpyxl

wb = openpyxl.load_workbook("excel/ssc_jsc_results_scraped_version_two.xlsx", data_only=True)
for sn in ('ssc', 'jsc'):
    ws = wb[sn]
    print(sn, "max_row=", ws.max_row, "max_col=", ws.max_column)
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    for row in ws.iter_rows(min_row=2, max_row=20, values_only=True):
        if row[idx['eiin']] == 108019:
            print(row)

# spot check 108019 across all ssc years
print("\n108019 ssc full year run:")
ws = wb['ssc']
header = [c.value for c in ws[1]]
idx = {h: i for i, h in enumerate(header)}
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[idx['eiin']] == 108019:
        print(row[idx['exam_year']], row[idx['examinee_total']], row[idx['sum_gpa']], row[idx['nonscience_examinee_total']])
