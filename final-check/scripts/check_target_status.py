import openpyxl

TARGETS = [(108019, y) for y in [2012,2013,2014,2015,2016,2017,2018,2019,2021,2022,2023,2024,2025,2026]] + [(108487, 2022)]

wb = openpyxl.load_workbook("../excel/ssc_jsc_results_scraped.xlsx", data_only=True)
ws = wb['ssc']
header = [c.value for c in ws[1]]
idx = {h: i for i, h in enumerate(header)}

for row in ws.iter_rows(min_row=2, values_only=True):
    key = (row[idx['eiin']], row[idx['exam_year']])
    if key in TARGETS:
        print(key, "status=", row[idx['scrape_status']], "note=", row[idx['scrape_note']],
              "examinee_total=", row[idx['examinee_total']])
