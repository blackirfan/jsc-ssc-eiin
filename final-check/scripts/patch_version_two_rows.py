import os, sys, importlib.util
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

spec = importlib.util.spec_from_file_location("bv2", os.path.join(ROOT, "scripts", "build_version_two.py"))
bv2 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bv2)

TARGETS = [('ssc', 108019, y) for y in [2018, 2019, 2021, 2022, 2023, 2024, 2025, 2026]]
# 108487/2022 appended only if its PDF exists yet
if os.path.exists(os.path.join(bv2.PDF_DIR, "ssc_108487_2022.pdf")):
    TARGETS.append(('ssc', 108487, 2022))

PATH = "excel/ssc_jsc_results_scraped_version_two.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)

for kind, eiin, year in TARGETS:
    ws = wb[kind]
    header = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(header)}
    data = bv2.compute_row(kind, eiin, year)
    if data is None:
        print(f"{kind} {eiin} {year}: still no PDF, skipping")
        continue
    row_num = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=idx['eiin']).value == eiin and ws.cell(row=r, column=idx['exam_year']).value == year:
            row_num = r
            break
    if row_num is None:
        print(f"{kind} {eiin} {year}: row not found!")
        continue
    for k, v in data.items():
        if k in idx:
            ws.cell(row=row_num, column=idx[k], value=v)
    if 'scrape_status' in idx:
        ws.cell(row=row_num, column=idx['scrape_status'], value='done')
    print(f"{kind} {eiin} {year}: patched -> examinee_total={data['examinee_total']}")

wb.save(PATH)
print("saved.")
