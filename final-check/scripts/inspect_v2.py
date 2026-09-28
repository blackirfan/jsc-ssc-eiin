import openpyxl

for fname in ["ssc_jsc_results_scraped_version_two.xlsx", "ssc_jsc_results_scraped_version_two_raw.xlsx"]:
    print("=====", fname, "=====")
    wb = openpyxl.load_workbook(f"excel/{fname}", data_only=True)
    print("sheets:", wb.sheetnames)
    for sn in wb.sheetnames:
        ws = wb[sn]
        print(f"--- {sn} --- max_row={ws.max_row} max_col={ws.max_column}")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True)):
            print(row)
    print()
