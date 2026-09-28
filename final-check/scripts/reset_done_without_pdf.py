import openpyxl
import os

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(SCRIPT_DIR, "..", ".."))
EXCEL_PATH = os.path.join(ROOT, "excel", "ssc_jsc_results_scraped.xlsx")
PDF_DIR = os.path.join(ROOT, "pdfs")

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

changed = 0
for kind in ('ssc', 'jsc'):
    ws = wb[kind]
    header = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(header)}
    for r in range(2, ws.max_row + 1):
        status = ws.cell(row=r, column=idx['scrape_status']).value
        if status != 'done':
            continue
        eiin = ws.cell(row=r, column=idx['eiin']).value
        year = ws.cell(row=r, column=idx['exam_year']).value
        pdf_path = os.path.join(PDF_DIR, f"{kind}_{eiin}_{year}.pdf")
        if not os.path.exists(pdf_path):
            print(f"{kind} eiin={eiin} year={year}: marked done but no PDF -> resetting to error")
            ws.cell(row=r, column=idx['scrape_status']).value = 'error'
            ws.cell(row=r, column=idx['scrape_note']).value = (
                'manual reset 2026-09-27: scrape_status was done but no PDF on disk '
                '(pre-fix scraper bug that marked rows done even when PDF download failed); '
                'queued for rescrape'
            )
            changed += 1

print("total rows reset:", changed)
if changed:
    wb.save(EXCEL_PATH)
    print("saved.")
else:
    print("nothing to save.")
