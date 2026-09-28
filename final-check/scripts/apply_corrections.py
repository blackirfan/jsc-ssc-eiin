import re, os, sys, glob, time, csv, shutil
import openpyxl
import importlib.util

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

spec = importlib.util.spec_from_file_location("audit", os.path.join(ROOT, "scripts", "audit.py"))
audit = importlib.util.module_from_spec(spec)
spec.loader.exec_module(audit)

SRC = "excel/ssc_jsc_results_scraped.xlsx"
DST = "excel/ssc_jsc_results_corrected.xlsx"
LOG_PATH = "audit_output/correction_log.csv"

CORRECTABLE_FIELDS = [
    'avg_gpa',
    'science_examinee_total', 'science_examinee_passed', 'science_avg_gpa',
    'nonscience_examinee_total', 'nonscience_examinee_passed', 'nonscience_avg_gpa',
]


def log(msg):
    ts = time.strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{ts}] {msg}")


def compute_pdf_values(kind, eiin, year):
    path = f"pdfs/{kind}_{eiin}_{year}.pdf"
    if not os.path.exists(path):
        return None
    parsed, err = audit.parse_pdf(path)
    if err:
        return None
    header = parsed['header']
    groups = parsed['groups']
    science_entries = groups.get('SCIENCE', [])
    nonscience_entries = []
    for gname, entries in groups.items():
        if gname != 'SCIENCE':
            nonscience_entries.extend(entries)
    all_entries = []
    for entries in groups.values():
        all_entries.extend(entries)

    overall_sum = audit.summarize(all_entries)
    science_sum = audit.summarize(science_entries)
    nonscience_sum = audit.summarize(nonscience_entries)

    overall_avg, _ = audit.avg_of(overall_sum)
    science_avg, _ = audit.avg_of(science_sum) if science_entries else (None, 'no_group')
    nonscience_avg, _ = audit.avg_of(nonscience_sum) if nonscience_entries else (None, 'no_group')

    science_gpa5_unreliable = science_sum['unknown_pass_count'] > 0
    nonscience_gpa5_unreliable = nonscience_sum['unknown_pass_count'] > 0

    return {
        'examinee_total': header['examinee'],
        'examinee_passed': header['passed'],
        'examinee_gpa5': header['gpa5'],
        'avg_gpa': overall_avg,
        'science_examinee_total': science_sum['total'] if science_entries else None,
        'science_examinee_passed': science_sum['passed'] if science_entries else None,
        'science_examinee_gpa5': None if science_gpa5_unreliable else (science_sum['gpa5'] if science_entries else None),
        'science_avg_gpa': science_avg,
        'nonscience_examinee_total': nonscience_sum['total'] if nonscience_entries else None,
        'nonscience_examinee_passed': nonscience_sum['passed'] if nonscience_entries else None,
        'nonscience_examinee_gpa5': None if nonscience_gpa5_unreliable else (nonscience_sum['gpa5'] if nonscience_entries else None),
        'nonscience_avg_gpa': nonscience_avg,
    }


def main():
    log(f"Copying {SRC} -> {DST} (original left untouched)")
    shutil.copyfile(SRC, DST)

    log("Opening copy for editing...")
    wb = openpyxl.load_workbook(DST, data_only=True)

    change_log = []
    stats = {'ssc': {'rows_checked': 0, 'rows_changed': 0, 'cells_changed': 0},
              'jsc': {'rows_checked': 0, 'rows_changed': 0, 'cells_changed': 0}}

    for kind in ('ssc', 'jsc'):
        ws = wb[kind]
        header_row = [c.value for c in ws[1]]
        col_idx = {h: i + 1 for i, h in enumerate(header_row)}  # 1-based for openpyxl cell access
        n = ws.max_row
        log(f"--- {kind}: scanning {n - 1} rows ---")
        for r in range(2, n + 1):
            eiin = ws.cell(row=r, column=col_idx['eiin']).value
            year = ws.cell(row=r, column=col_idx['exam_year']).value
            status = ws.cell(row=r, column=col_idx['scrape_status']).value
            if status != 'done' or eiin is None or year is None:
                continue
            stats[kind]['rows_checked'] += 1
            pdf_vals = compute_pdf_values(kind, int(eiin), int(year))
            if pdf_vals is None:
                continue  # missing/unparseable pdf, leave excel untouched

            row_changed = False
            for field in CORRECTABLE_FIELDS:
                pdf_val = pdf_vals[field]
                if field.endswith('avg_gpa') and pdf_val is None:
                    continue
                excel_cell = ws.cell(row=r, column=col_idx[field])
                excel_val = excel_cell.value
                cmp_excel, cmp_pdf = excel_val, pdf_val
                if not field.endswith('avg_gpa'):
                    if cmp_excel is None:
                        cmp_excel = 0
                    if cmp_pdf is None:
                        cmp_pdf = 0
                if not audit.close(cmp_excel, cmp_pdf):
                    change_log.append({
                        'kind': kind, 'eiin': eiin, 'year': year, 'field': field,
                        'old_value': excel_val, 'new_value': pdf_val,
                    })
                    excel_cell.value = pdf_val
                    row_changed = True
                    stats[kind]['cells_changed'] += 1
            if row_changed:
                stats[kind]['rows_changed'] += 1

            if stats[kind]['rows_checked'] % 1000 == 0:
                log(f"[{kind}] checked {stats[kind]['rows_checked']} rows, "
                    f"changed {stats[kind]['rows_changed']} rows / {stats[kind]['cells_changed']} cells so far")

        log(f"--- {kind} done: checked={stats[kind]['rows_checked']} "
            f"rows_changed={stats[kind]['rows_changed']} cells_changed={stats[kind]['cells_changed']} ---")

    log(f"Writing corrected workbook to {DST} ...")
    wb.save(DST)

    with open(LOG_PATH, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=['kind', 'eiin', 'year', 'field', 'old_value', 'new_value'])
        w.writeheader()
        w.writerows(change_log)

    total_rows_changed = stats['ssc']['rows_changed'] + stats['jsc']['rows_changed']
    total_cells_changed = stats['ssc']['cells_changed'] + stats['jsc']['cells_changed']
    log(f"=== DONE === total rows changed={total_rows_changed} total cells changed={total_cells_changed}")
    log(f"Change log written to {LOG_PATH} ({len(change_log)} entries)")


if __name__ == '__main__':
    main()
