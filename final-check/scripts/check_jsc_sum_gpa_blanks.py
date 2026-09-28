import openpyxl
from collections import Counter

wb = openpyxl.load_workbook("excel/ssc_jsc_results_scraped_version_two.xlsx", data_only=True)
ws = wb['jsc']
header = [c.value for c in ws[1]]
idx = {h: i for i, h in enumerate(header)}

blank_by_year = Counter()
done_by_year = Counter()
blank_but_done = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    status = row[idx['scrape_status']]
    year = row[idx['exam_year']]
    if status == 'done':
        done_by_year[year] += 1
        if row[idx['sum_gpa']] is None:
            blank_by_year[year] += 1
            blank_but_done += 1

print("done rows with blank sum_gpa, by year:")
for y in sorted(done_by_year):
    print(f"  {y}: blank={blank_by_year.get(y,0)} / done={done_by_year[y]}")
print("total done-but-blank-sum_gpa:", blank_but_done)
