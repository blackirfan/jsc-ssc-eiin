import os, json, shutil, importlib.util
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

spec = importlib.util.spec_from_file_location("bv2", os.path.join(ROOT, "scripts", "build_version_two.py"))
bv2 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bv2)

V2_PATH = "excel/ssc_jsc_results_scraped_version_two.xlsx"
V3_PATH = "excel/ssc_jsc_results_scraped_version_three.xlsx"
CHECKPOINT_PATH = "recheck_v3/checkpoint.json"

print("Copying", V2_PATH, "->", V3_PATH)
shutil.copyfile(V2_PATH, V3_PATH)

checkpoint = json.load(open(CHECKPOINT_PATH, encoding="utf-8"))


def recheck_outcome(kind, eiin, year):
    key = f"{eiin}|{kind}|{year}"
    entry = checkpoint.get(key)
    if entry is None:
        return None  # wasn't a no_result row, not rechecked
    if entry.get('found'):
        return 'found'
    if entry.get('noResult'):
        return 'confirmed_no_result'
    return 'error_inconclusive'


wb = openpyxl.load_workbook(V3_PATH, data_only=True)

for kind in ('ssc', 'jsc'):
    ws = wb[kind]
    header = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(header)}

    col_name = 'recheck_2026_10_03'
    if col_name not in header:
        new_col = len(header) + 1
        ws.cell(row=1, column=new_col, value=col_name)
        idx[col_name] = new_col
        header.append(col_name)

    counts = {'found': 0, 'confirmed_no_result': 0, 'error_inconclusive': 0, 'not_rechecked': 0}
    newly_filled = 0
    for r in range(2, ws.max_row + 1):
        eiin = ws.cell(row=r, column=idx['eiin']).value
        year = ws.cell(row=r, column=idx['exam_year']).value
        outcome = recheck_outcome(kind, eiin, year)
        if outcome is None:
            ws.cell(row=r, column=idx[col_name], value='not_rechecked')
            counts['not_rechecked'] += 1
            continue
        ws.cell(row=r, column=idx[col_name], value=outcome)
        counts[outcome] += 1

        if outcome == 'found':
            # already patched into version_two (and thus copied into version_three)
            # via patch_version_two_rows.py, but recompute here too for safety/idempotency
            if ws.cell(row=r, column=idx['examinee_total']).value is None:
                data = bv2.compute_row(kind, eiin, year)
                if data:
                    for k, v in data.items():
                        if k in idx:
                            ws.cell(row=r, column=idx[k], value=v)
                    if 'scrape_status' in idx:
                        ws.cell(row=r, column=idx['scrape_status'], value='done')
                    newly_filled += 1

    print(kind, counts, "newly_filled_here:", newly_filled)

wb.save(V3_PATH)
print("Saved:", V3_PATH)
