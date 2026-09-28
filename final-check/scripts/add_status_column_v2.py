import openpyxl

OLD_PATH = "../excel/ssc_jsc_results_scraped.xlsx"
NEW_PATH = "excel/ssc_jsc_results_scraped_version_two.xlsx"

import os
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OLD_PATH = os.path.normpath(os.path.join(SCRIPT_DIR, "..", "..", "excel", "ssc_jsc_results_scraped.xlsx"))
NEW_PATH = os.path.normpath(os.path.join(SCRIPT_DIR, "..", "excel", "ssc_jsc_results_scraped_version_two.xlsx"))

print("Old (status source):", OLD_PATH)
print("New (target):", NEW_PATH)

old_wb = openpyxl.load_workbook(OLD_PATH, data_only=True)
new_wb = openpyxl.load_workbook(NEW_PATH, data_only=True)

for sn in ('ssc', 'jsc'):
    old_ws = old_wb[sn]
    old_header = [c.value for c in old_ws[1]]
    old_idx = {h: i for i, h in enumerate(old_header)}
    status_lookup = {}
    for row in old_ws.iter_rows(min_row=2, values_only=True):
        key = (row[old_idx['eiin']], row[old_idx['exam_year']])
        status_lookup[key] = row[old_idx['scrape_status']]

    new_ws = new_wb[sn]
    header = [c.value for c in new_ws[1]]
    if 'scrape_status' not in header:
        status_col = len(header) + 1
        new_ws.cell(row=1, column=status_col, value='scrape_status')
    else:
        status_col = header.index('scrape_status') + 1
    idx = {h: i + 1 for i, h in enumerate(header)}

    counts = {'done': 0, 'no_result': 0}
    for r in range(2, new_ws.max_row + 1):
        eiin = new_ws.cell(row=r, column=idx['eiin']).value
        year = new_ws.cell(row=r, column=idx['exam_year']).value
        old_status = status_lookup.get((eiin, year))
        status = 'done' if old_status == 'done' else 'no_result'
        new_ws.cell(row=r, column=status_col, value=status)
        counts[status] += 1

    print(sn, "status column written:", counts)

new_wb.save(NEW_PATH)
print("saved.")
