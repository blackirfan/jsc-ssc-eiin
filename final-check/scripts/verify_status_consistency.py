import openpyxl

wb = openpyxl.load_workbook("excel/ssc_jsc_results_scraped_version_two.xlsx", data_only=True)
for sn in ('ssc', 'jsc'):
    ws = wb[sn]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    mismatches = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        has_data = row[idx['examinee_total']] is not None
        status = row[idx['scrape_status']]
        if has_data and status != 'done':
            mismatches += 1
        if not has_data and status == 'done':
            mismatches += 1
    print(sn, "status/data inconsistencies:", mismatches)
