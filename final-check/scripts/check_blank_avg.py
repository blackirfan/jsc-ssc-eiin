import os, sys, importlib.util, glob, re

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

spec = importlib.util.spec_from_file_location("audit", os.path.join(ROOT, "scripts", "audit.py"))
audit = importlib.util.module_from_spec(spec)
spec.loader.exec_module(audit)

import openpyxl
wb = openpyxl.load_workbook("excel/ssc_jsc_results_corrected.xlsx", data_only=True)

for sn in ('ssc', 'jsc'):
    ws = wb[sn]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    blank_avg = 0
    blank_science_avg = 0
    blank_nonscience_avg = 0
    recoverable_avg = 0
    recoverable_science = 0
    recoverable_nonscience = 0
    examples = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx['scrape_status']] != 'done':
            continue
        eiin = row[idx['eiin']]
        year = row[idx['exam_year']]
        av = row[idx['avg_gpa']]
        sav = row[idx['science_avg_gpa']]
        nav = row[idx['nonscience_avg_gpa']]
        need_check = (av is None) or (nav is None and row[idx['nonscience_examinee_total']] not in (None, 0)) or (sav is None and row[idx['science_examinee_total']] not in (None, 0))
        if av is None:
            blank_avg += 1
        if nav is None and row[idx['nonscience_examinee_total']] not in (None, 0):
            blank_nonscience_avg += 1
        if sav is None and row[idx['science_examinee_total']] not in (None, 0):
            blank_science_avg += 1
        if need_check:
            path = f"pdfs/{sn}_{eiin}_{year}.pdf"
            if not os.path.exists(path):
                continue
            parsed, err = audit.parse_pdf(path)
            if err:
                continue
            groups = parsed['groups']
            science_entries = groups.get('SCIENCE', [])
            nonscience_entries = []
            for gname, entries in groups.items():
                if gname != 'SCIENCE':
                    nonscience_entries.extend(entries)
            all_entries = []
            for entries in groups.values():
                all_entries.extend(entries)
            overall_sum = audit.summarize(all_entries)
            science_sum = audit.summarize(science_entries)
            nonscience_sum = audit.summarize(nonscience_entries)
            oa, os_ = audit.avg_of(overall_sum)
            sa, ss_ = audit.avg_of(science_sum) if science_entries else (None, 'no_group')
            na, ns_ = audit.avg_of(nonscience_sum) if nonscience_entries else (None, 'no_group')
            if av is None and oa is not None:
                recoverable_avg += 1
                if len(examples) < 5:
                    examples.append((sn, eiin, year, 'avg_gpa', av, oa, os_))
            if nav is None and row[idx['nonscience_examinee_total']] not in (None, 0) and na is not None:
                recoverable_nonscience += 1
                if len(examples) < 10:
                    examples.append((sn, eiin, year, 'nonscience_avg_gpa', nav, na, ns_))
            if sav is None and row[idx['science_examinee_total']] not in (None, 0) and sa is not None:
                recoverable_science += 1
                if len(examples) < 15:
                    examples.append((sn, eiin, year, 'science_avg_gpa', sav, sa, ss_))

    print(f"=== {sn} ===")
    print("blank avg_gpa (done rows):", blank_avg)
    print("blank nonscience_avg_gpa where nonscience group nonzero:", blank_nonscience_avg)
    print("blank science_avg_gpa where science group nonzero:", blank_science_avg)
    print("RECOVERABLE avg_gpa (pdf has a real answer):", recoverable_avg)
    print("RECOVERABLE nonscience_avg_gpa:", recoverable_nonscience)
    print("RECOVERABLE science_avg_gpa:", recoverable_science)
    print("examples:")
    for e in examples:
        print("  ", e)
    print()
