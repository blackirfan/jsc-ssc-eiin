import openpyxl
import os

TARGETS = [
    (108019, y) for y in [2012, 2013, 2014, 2015, 2016, 2017, 2018, 2019, 2021, 2022, 2023, 2024, 2025, 2026]
] + [(108487, 2022)]

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PATH = os.path.normpath(os.path.join(SCRIPT_DIR, "..", "..", "excel", "ssc_jsc_results_scraped.xlsx"))
print("Using excel file:", PATH)

wb = openpyxl.load_workbook(PATH, data_only=True)
ws = wb['ssc']
header = [c.value for c in ws[1]]
idx = {h: i + 1 for i, h in enumerate(header)}

changed = 0
for r in range(2, ws.max_row + 1):
    eiin = ws.cell(row=r, column=idx['eiin']).value
    year = ws.cell(row=r, column=idx['exam_year']).value
    if (eiin, year) in TARGETS:
        status_cell = ws.cell(row=r, column=idx['scrape_status'])
        note_cell = ws.cell(row=r, column=idx['scrape_note'])
        print(f"row {r}: eiin={eiin} year={year} old_status={status_cell.value!r}")
        status_cell.value = 'error'
        note_cell.value = 'manual reset 2026-09-27: PDF missing on disk despite done status, flagged by final-check audit; queued for rescrape'
        changed += 1

print("changed rows:", changed, "of", len(TARGETS), "targets")
wb.save(PATH)
print("saved.")
