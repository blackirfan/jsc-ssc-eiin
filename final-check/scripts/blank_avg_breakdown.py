import openpyxl
from collections import Counter

wb = openpyxl.load_workbook("excel/ssc_jsc_results_corrected.xlsx", data_only=True)
for sn in ('ssc', 'jsc'):
    ws = wb[sn]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    c = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx['scrape_status']] != 'done':
            continue
        if row[idx['avg_gpa']] is None:
            c[row[idx['exam_year']]] += 1
    print(sn, "blank avg_gpa by year:", dict(sorted(c.items())))
